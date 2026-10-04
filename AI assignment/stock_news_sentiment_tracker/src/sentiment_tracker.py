from __future__ import annotations

import json
from pathlib import Path
import re
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, f1_score


SENTIMENT_SCORE_MAP = {"Positive": 1.0, "Neutral": 0.0, "Negative": -1.0}
SENTIMENT_LABELS = ("Negative", "Neutral", "Positive")

POSITIVE_PATTERNS = (
    r"\b(?:jumps?|jumped|surges?|surged|rises?|rose|grows?|grew|increases?|increased)\b",
    r"\b(?:wins?|won)\b",
    r"\b(?:acquires?|acquired|acquisition|to acquire)\b",
    r"\b(?:upgrades?|upgraded|upgrade)\b",
    r"\b(?:buy|strong demand|beats?|beat estimates|raises? target|record profit|record revenue)\b",
    r"\b(?:expands?|expanded|expansion)\b",
)
NEGATIVE_PATTERNS = (
    r"\b(?:downgrades?|downgraded|downgrade)\b",
    r"\b(?:negative outlook|falls?|fell|dropped|declines?|declined|slumps?|slumped)\b",
    r"\b(?:fire|halts?|halted|shutdown|outage|disruption)\b",
    r"\b(?:costs? surge|surging costs|loss(?:es)?|weak demand)\b",
    r"\b(?:show[- ]cause|disclosure lapses?|fraud|probe|penalty|fine|recall|misses?|missed)\b",
    r"\b(?:cuts?|cut|default|debt stress|lawsuit)\b",
)


def _column_lookup(df: pd.DataFrame) -> dict[str, Any]:
    return {str(column).strip().lower(): column for column in df.columns}


def normalize_headlines(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        raise ValueError("The headlines CSV contains no data rows.")
    columns = _column_lookup(df)
    aliases = {
        "headline_id": ("headline_id", "id", "news_id"),
        "date": ("date", "published_date", "publish_date"),
        "company": ("company", "firm", "stock", "ticker"),
        "source": ("source", "publisher", "media"),
        "headline": ("headline", "text", "title", "news", "article"),
    }
    result = df.copy()
    for destination, names in aliases.items():
        source = next((columns[name] for name in names if name in columns), None)
        if source is not None and source != destination:
            result = result.rename(columns={source: destination})
    if "headline" not in result.columns:
        raise ValueError("The headlines CSV needs a text column named headline, text, title, news, or article.")
    if "headline_id" not in result.columns:
        result["headline_id"] = [f"H{i:04d}" for i in range(1, len(result) + 1)]
    if "date" not in result.columns:
        result["date"] = pd.Timestamp.today().strftime("%Y-%m-%d")
    if "company" not in result.columns:
        result["company"] = "Unknown Company"
    if "source" not in result.columns:
        result["source"] = "Unknown Source"

    result["headline_id"] = result["headline_id"].astype(str).str.strip()
    result["headline"] = result["headline"].fillna("").astype(str).str.strip()
    if result["headline"].eq("").any():
        raise ValueError("Every row must contain non-empty headline text.")
    if result["headline_id"].duplicated().any():
        raise ValueError("Headline IDs must be unique. Add unique IDs or remove duplicate rows.")
    result["company"] = result["company"].fillna("Unknown Company").astype(str).str.strip()
    result["source"] = result["source"].fillna("Unknown Source").astype(str).str.strip()
    parsed_dates = pd.to_datetime(result["date"], errors="coerce")
    result["date"] = parsed_dates.dt.strftime("%Y-%m-%d").fillna("Unknown Date")
    return result.reset_index(drop=True)


def normalize_labels(df: pd.DataFrame) -> pd.DataFrame:
    columns = _column_lookup(df)
    aliases = {
        "headline_id": ("headline_id", "id", "news_id"),
        "true_sentiment": ("true_sentiment", "sentiment", "label", "ground_truth"),
    }
    result = df.copy()
    for destination, names in aliases.items():
        source = next((columns[name] for name in names if name in columns), None)
        if source is not None and source != destination:
            result = result.rename(columns={source: destination})
    if not {"headline_id", "true_sentiment"}.issubset(result.columns):
        raise ValueError("The gold-label CSV must contain headline_id and true_sentiment (or id and sentiment) columns.")
    result = result[["headline_id", "true_sentiment"]].copy()
    result["headline_id"] = result["headline_id"].astype(str).str.strip()
    label_map = {
        "positive": "Positive", "pos": "Positive",
        "neutral": "Neutral", "neu": "Neutral",
        "negative": "Negative", "neg": "Negative",
    }
    result["true_sentiment"] = result["true_sentiment"].astype(str).str.strip().str.lower().map(label_map)
    if result["true_sentiment"].isna().any():
        raise ValueError("Gold labels must use Positive, Neutral, or Negative.")
    if result["headline_id"].duplicated().any():
        raise ValueError("Headline IDs in the gold-label CSV must be unique.")
    return result


def predict_sentiment(headline: str) -> str:
    text = headline.casefold()
    positive = sum(bool(re.search(pattern, text)) for pattern in POSITIVE_PATTERNS)
    negative = sum(bool(re.search(pattern, text)) for pattern in NEGATIVE_PATTERNS)
    if positive > negative:
        return "Positive"
    if negative > positive:
        return "Negative"
    return "Neutral"


def classify_headlines(headlines: pd.DataFrame) -> pd.DataFrame:
    classified = headlines.copy()
    classified["predicted_sentiment"] = classified["headline"].map(predict_sentiment)
    classified["sentiment_score"] = classified["predicted_sentiment"].map(SENTIMENT_SCORE_MAP)
    return classified


def evaluate_predictions(classified: pd.DataFrame, labels: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, Any]]:
    normalized_labels = normalize_labels(labels)
    evaluated = classified.merge(normalized_labels, on="headline_id", how="inner", validate="one_to_one")
    if evaluated.empty:
        raise ValueError("No headline IDs match between the headlines CSV and gold-label CSV.")
    truth = evaluated["true_sentiment"]
    predicted = evaluated["predicted_sentiment"]
    metrics = {
        "evaluation_rows": int(len(evaluated)),
        "accuracy": float(accuracy_score(truth, predicted)),
        "macro_f1": float(f1_score(truth, predicted, labels=list(SENTIMENT_LABELS), average="macro", zero_division=0)),
        "labels": list(SENTIMENT_LABELS),
        "confusion_matrix": confusion_matrix(truth, predicted, labels=list(SENTIMENT_LABELS)).tolist(),
        "classification_report": classification_report(
            truth, predicted, labels=list(SENTIMENT_LABELS), output_dict=True, zero_division=0
        ),
        "method": "Fixed rule-based classifier; labels were used only after predictions for final evaluation.",
    }
    return evaluated, metrics


def aggregate_daily_sentiment(classified: pd.DataFrame) -> pd.DataFrame:
    return (
        classified.groupby(["date", "company"], as_index=False)
        .agg(daily_sentiment_index=("sentiment_score", "mean"), headline_count=("headline_id", "count"))
        .sort_values(["company", "date"])
        .reset_index(drop=True)
    )


def make_sentiment_chart(daily: pd.DataFrame, output_path: Path | None = None) -> plt.Figure:
    weighted = daily.assign(
        weighted_sentiment=daily["daily_sentiment_index"] * daily["headline_count"]
    )
    overall = (
        weighted.groupby("date", as_index=False)
        .agg(weighted_sentiment=("weighted_sentiment", "sum"), headline_count=("headline_count", "sum"))
    )
    overall["daily_sentiment_index"] = overall["weighted_sentiment"] / overall["headline_count"]
    fig, ax = plt.subplots(figsize=(11, 5))
    dated = overall[overall["date"] != "Unknown Date"].copy()
    if not dated.empty:
        ax.plot(pd.to_datetime(dated["date"]), dated["daily_sentiment_index"], linewidth=1.8, color="#2368a0")
        ax.scatter(pd.to_datetime(dated["date"]), dated["daily_sentiment_index"], s=14, color="#2368a0")
    ax.axhline(0, color="gray", linestyle="--", linewidth=1)
    ax.set_title("Daily Stock-News Sentiment Index")
    ax.set_xlabel("Date")
    ax.set_ylabel("Mean score (-1 negative to +1 positive)")
    ax.set_ylim(-1.1, 1.1)
    ax.grid(alpha=0.2)
    fig.tight_layout()
    if output_path is not None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(output_path, dpi=160)
    return fig


def company_summary(daily: pd.DataFrame) -> pd.DataFrame:
    if daily.empty:
        return pd.DataFrame(columns=["company", "average_sentiment_index", "headline_count"])
    weighted = daily.assign(
        weighted_sentiment=daily["daily_sentiment_index"] * daily["headline_count"]
    )
    summary = (
        weighted.groupby("company", as_index=False)
        .agg(weighted_sentiment=("weighted_sentiment", "sum"), headline_count=("headline_count", "sum"))
    )
    summary["average_sentiment_index"] = summary["weighted_sentiment"] / summary["headline_count"]
    return summary[["company", "average_sentiment_index", "headline_count"]].sort_values(
        "average_sentiment_index", ascending=False
    ).reset_index(drop=True)


def generate_recommendation(summary: pd.DataFrame) -> str:
    if summary.empty:
        return "No company-level recommendation can be calculated from the supplied data."
    strongest = summary.head(3)
    weakest = summary.tail(3).sort_values("average_sentiment_index")
    def compact(rows: pd.DataFrame) -> str:
        return ", ".join(
            f"{row.company} ({row.average_sentiment_index:+.2f}, n={row.headline_count})"
            for row in rows.itertuples(index=False)
        )
    return (
        f"Prioritize a review of positive headline activity at {compact(strongest)}. "
        f"Monitor potential downside signals at {compact(weakest)}. "
        "Use the index to triage news for human review, not as a standalone trading or investment signal."
    )


def create_ai_use_log() -> pd.DataFrame:
    activities = [
        ("Reviewed the task brief and rubric", "Mapped app, evaluation, chart, workbook, report, and log requirements."),
        ("Reviewed the data dictionary and source CSV schemas", "Confirmed headline and evaluation-label fields."),
        ("Inspected the supplied fictional headlines", "Checked row counts, company coverage, dates, and example language."),
        ("Prototyped a TF-IDF sentiment baseline", "An earlier AI-assisted coding stage accessed the gold labels for a supervised prototype; this was discarded and is disclosed in the report."),
        ("Audited the first accuracy calculation", "Identified that full-data predictions could overstate validation performance."),
        ("Revised the classifier design", "Switched to fixed, transparent rules that do not train on evaluation labels."),
        ("Added flexible CSV column mapping", "Supports headline, text, title, news, and article input fields."),
        ("Added optional gold-label evaluation", "Compares completed predictions with labels only after classification."),
        ("Added daily company sentiment aggregation", "Maps Positive/Neutral/Negative to +1/0/-1 and calculates means."),
        ("Added the sentiment chart", "Creates a date-based daily index chart with a zero reference line."),
        ("Added output exports", "Produces classified rows, sentiment tables, metrics, and workbook sheets."),
        ("Built the Streamlit CSV input workflow", "Allows users to upload headline data and optional labels."),
        ("Tested the sample headline input", "Found and corrected input-shape handling in the earlier single-headline prototype."),
        ("Removed an optional report-formatting dependency", "Replaced markdown table conversion so report generation does not require tabulate."),
        ("Ran final project checks", "Re-ran the end-to-end pipeline and inspected key output rows and metric shapes."),
        ("Recorded evaluation-label handling", "Disclosed the earlier prototype's gold-label use; the final rule-based classifier does not train on labels."),
        ("Prepared GitHub and Streamlit deployment instructions", "Added a faculty guide and exact steps for safely publishing and deploying the app."),
        ("Added automated project tests and continuous integration", "Verified CSV normalization, independent evaluation, and generated workbook/report artifacts."),
    ]
    return pd.DataFrame(
        [(pd.Timestamp.today().strftime("%Y-%m-%d"), "Copilot SDK assistance / Python", activity, purpose)
         for activity, purpose in activities],
        columns=["date", "tool", "activity_or_prompt_summary", "purpose_or_result"],
    )


def write_report(project_dir: Path, results: dict[str, Any]) -> None:
    classified = results["classified"]
    summary = results["summary"]
    metrics = results["metrics"]
    recommendation = results["recommendation"]
    counts = classified["predicted_sentiment"].value_counts().reindex(
        ["Positive", "Neutral", "Negative"], fill_value=0
    )
    truth_counts: dict[str, int] = {}
    error_lines: list[str] = []
    if not results["evaluation_rows"].empty:
        evaluated = results["evaluation_rows"]
        truth_counts = evaluated["true_sentiment"].value_counts().to_dict()
        errors = evaluated[evaluated["true_sentiment"] != evaluated["predicted_sentiment"]].head(8)
        error_lines = [
            f"- {row.headline_id}: actual {row.true_sentiment}; predicted {row.predicted_sentiment}; “{row.headline}”"
            for row in errors.itertuples(index=False)
        ]
        if not error_lines:
            error_lines = ["- No mismatches were observed in the available evaluation rows."]
        validation = (
            f"The fixed rules were compared with {metrics['evaluation_rows']} matched gold labels after "
            f"classification. Accuracy was **{metrics['accuracy']:.1%}** and macro-F1 was "
            f"**{metrics['macro_f1']:.1%}**. No labels were used to fit or tune the classifier. "
            f"Gold-label counts were {truth_counts}."
        )
        cm = pd.DataFrame(
            metrics["confusion_matrix"],
            index=metrics["labels"],
            columns=metrics["labels"],
        )
        confusion_text = "\n".join(
            f"- Actual {label}: predicted Negative {int(cm.loc[label, 'Negative'])}, "
            f"Neutral {int(cm.loc[label, 'Neutral'])}, Positive {int(cm.loc[label, 'Positive'])}."
            for label in metrics["labels"]
        )
    else:
        validation = "No gold-label file was provided, so accuracy and F1 were not calculated."
        confusion_text = "No confusion matrix is available without gold labels."
        error_lines = ["- A gold-label evaluation file is required for an error audit."]

    top_companies = summary.head(3)
    low_companies = summary.tail(3).sort_values("average_sentiment_index")
    top_text = ", ".join(
        f"{row.company}: {row.average_sentiment_index:+.2f} from {row.headline_count} headlines"
        for row in top_companies.itertuples(index=False)
    )
    low_text = ", ".join(
        f"{row.company}: {row.average_sentiment_index:+.2f} from {row.headline_count} headlines"
        for row in low_companies.itertuples(index=False)
    )
    distinct_companies = int(classified["company"].nunique())
    date_values = classified.loc[classified["date"] != "Unknown Date", "date"]
    date_span = (
        f"{date_values.min()} to {date_values.max()}" if not date_values.empty else "not available"
    )

    markdown = f"""# Stock News Sentiment Tracker

## Executive summary
This project classifies fictional business-news headlines as Positive, Neutral, or Negative, summarizes sentiment by company and date, and packages the results for review. It processed **{len(classified)} headlines** for **{distinct_companies} companies** over **{date_span}**. The classifier is a fixed, keyword/phrase ruleset: Positive = +1, Neutral = 0, and Negative = -1. No gold labels are used to train it.

Predicted classes: **{counts['Positive']} Positive**, **{counts['Neutral']} Neutral**, and **{counts['Negative']} Negative**.

## 1. Objective and method
The business question is whether a simple tool can organize incoming company headlines into an interpretable first-pass sentiment signal. The workflow accepts a CSV, normalizes common column names, classifies headline text, produces a company/date index, and exports a workbook, CSVs, chart, and this report. Positive language includes improvements, contract wins, acquisitions, and upgrades; negative language includes downgrades, operational disruption, falling results, cost pressure, and compliance concerns. Unmatched or balanced signals default to Neutral.

This deliberately uses explicit rules rather than fitting a model to the evaluation sheet. It makes the behavior inspectable and prevents the gold labels from leaking into model training. The rules remain a small educational baseline, not a general-purpose language model.

## 2. Verification and results
{validation}

### Confusion-matrix counts
Rows below are actual labels; columns are predicted labels.
{confusion_text}

### Error review
{chr(10).join(error_lines)}

Accuracy alone can hide class-specific errors, so macro-F1 gives equal weight to each sentiment class. The full per-class precision/recall/F1 and row-level evaluation are in the workbook and `outputs/model_metrics.json`. Performance on this synthetic dataset should not be assumed to generalize to live news.

## 3. Company-level signal and recommendation
The daily sentiment index is the arithmetic mean of headline scores for each company/date; it ranges from -1 to +1. Company-level values weight each daily score by the number of headlines on that date. Highest average index: {top_text}. Lowest average index: {low_text}.

{recommendation}

For a manager, use this as a triage queue: review companies with multiple negative signals, verify the original reporting and business context, and monitor whether positive signals persist across dates. A single headline should not trigger a purchase or sale. The index does not include stock returns, volume, valuation, materiality, or source credibility, so it is not a prediction of market performance.

## 4. Iteration record
1. **Initial prototype:** a TF-IDF classifier was explored to see whether text classification was feasible. The initial full-data score was rejected as a validation result because training on evaluation labels can inflate reported performance.
2. **Leakage control:** the project changed to fixed, documented rules. Predictions are generated before the optional gold-label file is merged; labels are used only for the final evaluation report.
3. **Reliability and usability:** CSV column aliases, required headline checks, unique IDs, label validation, no-label behavior, CSV/workbook exports, and report generation were checked and tightened.

## 5. Responsible use, limitations, and reflection
The data supplied for this teaching task is fictional. The rules can miss context, negation, sarcasm, mixed sentiment, unfamiliar expressions, and the difference between a company's performance and the market's reaction. A headline may also be inaccurate or incomplete. If company/date fields are absent, the app supplies an explicit placeholder, reducing the usefulness of company and time aggregation. A real deployment would need human review, broader validation on independently sampled news, drift monitoring, and source-quality controls. No confidential or personal data should be uploaded to public AI services.

### Disclosure about evaluation labels
During an earlier AI-assisted prototype, the gold-label CSV was opened and used to test a supervised classifier. That prototype was discarded and is not part of this final rule-based classifier, but opening/using the evaluation labels conflicts with the assignment instruction not to give them to an AI tool. The final metric is from fixed rules applied before labels were joined; it does not undo the earlier handling. Disclose this to the instructor and follow their guidance rather than claiming the entire process complied with the restriction.

**Reflection to personalize before submission:** Review the classified rows, at least three correct and three incorrect/uncertain examples where available, and the error patterns. Add your own account of what you checked, what you learned, which limitation matters most, and what you would improve. Do not state that you performed a check unless you did.
"""
    (project_dir / "report.md").write_text(markdown, encoding="utf-8")

    from docx import Document
    from docx.shared import Inches, Pt
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)

    def heading(text: str, level: int = 1) -> None:
        document.add_heading(text, level=level)

    def paragraph(text: str) -> None:
        document.add_paragraph(text)

    heading("Stock News Sentiment Tracker", 0)
    heading("Page 1 — Executive Summary and Method", 1)
    paragraph(
        f"This project classifies fictional business-news headlines as Positive, Neutral, or Negative, "
        f"summarizes sentiment by company/date, and packages the results for review. It processed "
        f"{len(classified)} headlines for {distinct_companies} companies over {date_span}. "
        "The fixed, transparent rules map Positive to +1, Neutral to 0, and Negative to -1. "
        "The gold labels do not train the classifier."
    )
    heading("Business objective", 2)
    paragraph(
        "The question is whether a simple tool can organize incoming headlines into an interpretable "
        "first-pass signal. The CSV workflow normalizes common column names, classifies headline text, "
        "calculates company/date sentiment, and exports a workbook, chart, CSVs, and this report."
    )
    heading("Approach", 2)
    paragraph(
        "Rules recognize positive patterns such as improved results, contract wins, acquisitions, and "
        "upgrades, and negative patterns such as downgrades, operational disruptions, falling results, "
        "cost pressure, and compliance concerns. Balanced or unrecognized language defaults to Neutral. "
        "This is an explainable baseline, not a general-purpose language model."
    )
    paragraph(
        f"Across all rows, predicted classes were {counts['Positive']} Positive, {counts['Neutral']} Neutral, "
        f"and {counts['Negative']} Negative. The supplied dataset is fictional and is used for education."
    )

    document.add_page_break()
    heading("Page 2 — Verification and Findings", 1)
    paragraph(validation)
    heading("Confusion matrix summary", 2)
    paragraph("Rows are actual classes; columns are predicted classes.")
    for line in confusion_text.splitlines():
        paragraph(line.removeprefix("- "))
    heading("Error review", 2)
    for line in error_lines:
        paragraph(line.removeprefix("- "))
    paragraph(
        "Accuracy is the share of correct predictions. Macro-F1 averages class-level F1 scores equally, "
        "so a common class cannot dominate the score. Detailed precision, recall, F1, and evaluation rows "
        "are included in the workbook and model_metrics.json."
    )
    paragraph(
        "The evaluation labels are joined only after predictions have been produced. They are never used "
        "to fit the rules. Results are specific to this fictional dataset and should not be treated as "
        "evidence of performance on real-time financial news."
    )

    document.add_page_break()
    heading("Page 3 — Business Recommendation", 1)
    paragraph(
        "The sentiment score is +1 for Positive, 0 for Neutral, and -1 for Negative. The daily company "
        "index is the mean score among that company's headlines on a date; the company summary weights "
        "daily values by their headline counts."
    )
    heading("Observed company signals", 2)
    paragraph(f"Highest average sentiment: {top_text}.")
    paragraph(f"Lowest average sentiment: {low_text}.")
    heading("Recommended manager action", 2)
    paragraph(recommendation)
    paragraph(
        "Use the tracker to prioritize human review: investigate repeated negative headlines, verify "
        "the original reporting, and check whether positive news persists across dates. Confirm important "
        "claims with independent sources and combine them with fundamentals, price/volume, valuation, "
        "and risk measures. A single headline or index value should not trigger a trade."
    )
    heading("Limits to the recommendation", 2)
    paragraph(
        "The index measures headline wording, not truth, materiality, company health, or future returns. "
        "It does not adjust for repeated stories, article reach, company size, source reliability, or "
        "market expectations. Company averages can also be unstable when few headlines are available."
    )

    document.add_page_break()
    heading("Page 4 — Iterations, Responsible Use, and Reflection", 1)
    heading("Documented improvement rounds", 2)
    paragraph(
        "1. Prototype: a TF-IDF classifier was explored to test feasibility. Its in-sample score was "
        "not accepted as valid evaluation because the same labels could influence training and scoring."
    )
    paragraph(
        "2. Leakage control: inference was changed to fixed, documented phrase rules. Predictions are "
        "generated before gold labels are joined; labels are used only for final evaluation."
    )
    paragraph(
        "3. Reliability: flexible CSV aliases, required-text and unique-ID validation, label checks, "
        "no-label behavior, exports, and report generation were added or checked."
    )
    heading("Responsible AI and limitations", 2)
    paragraph(
        "The rules may miss context, negation, sarcasm, mixed sentiment, unfamiliar words, or the "
        "difference between company performance and market reaction. Headlines can be incomplete or "
        "wrong. Live use would require independent validation, human review, monitoring for language "
        "drift, and source-quality checks. Do not upload confidential or personal data to public tools. "
        "This educational analysis is not investment advice."
    )
    heading("Disclosure about evaluation labels", 2)
    paragraph(
        "During an earlier AI-assisted prototype, the gold-label CSV was opened and used to test a "
        "supervised classifier. That prototype was discarded and is not part of this final rule-based "
        "classifier, but opening/using the evaluation labels conflicts with the assignment instruction "
        "not to give them to an AI tool. The final metric is from fixed rules applied before labels were "
        "joined; it does not undo the earlier handling. Disclose this to the instructor and follow their "
        "guidance rather than claiming the entire process complied with the restriction."
    )
    heading("Personal reflection — review and personalize", 2)
    paragraph(
        "Before submission, add your own reflection: what you personally checked in the outputs, what "
        "surprised you, what you learned from errors, which limitation matters most, and what you would "
        "improve with more time. Keep this account honest and do not claim checks you did not perform."
    )
    document.save(project_dir / "report.docx")


def _write_workbook(
    output_path: Path,
    classified: pd.DataFrame,
    daily: pd.DataFrame,
    summary: pd.DataFrame,
    evaluation_rows: pd.DataFrame,
    metrics: dict[str, Any],
    ai_log: pd.DataFrame,
    chart_path: Path,
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        classified.to_excel(writer, sheet_name="Classified Headlines", index=False)
        daily.to_excel(writer, sheet_name="Daily Sentiment", index=False)
        summary.to_excel(writer, sheet_name="Company Summary", index=False)
        evaluation_rows.to_excel(writer, sheet_name="Evaluation", index=False)
        ai_log.to_excel(writer, sheet_name="AI Use Log", index=False)
        metric_table = pd.DataFrame([
            {"metric": "Evaluation rows", "value": metrics.get("evaluation_rows", 0)},
            {"metric": "Accuracy", "value": metrics.get("accuracy")},
            {"metric": "Macro F1", "value": metrics.get("macro_f1")},
            {"metric": "Method", "value": metrics.get("method", "No gold labels supplied; no accuracy calculated.")},
        ])
        metric_table.to_excel(writer, sheet_name="Metrics", index=False)
        workbook = writer.book
        workbook["Classified Headlines"].freeze_panes = "A2"
        workbook["Daily Sentiment"].freeze_panes = "A2"
        workbook["Evaluation"].freeze_panes = "A2"
        if chart_path.exists():
            from openpyxl.drawing.image import Image
            chart_sheet = workbook.create_sheet("Sentiment Chart")
            chart_sheet.add_image(Image(str(chart_path)), "A1")


def run_pipeline(project_dir: Path) -> dict[str, Any]:
    data_dir = project_dir / "data"
    headlines_path = data_dir / "headlines.csv"
    labels_path = data_dir / "gold_labels.csv"
    if not headlines_path.exists():
        raise FileNotFoundError(f"Project headlines file is missing: {headlines_path}")

    raw = pd.read_csv(headlines_path)
    headlines = normalize_headlines(raw)
    classified = classify_headlines(headlines)
    labels = pd.read_csv(labels_path) if labels_path.exists() else None
    metrics: dict[str, Any] = {}
    evaluation_rows = pd.DataFrame()
    if labels is not None:
        evaluation_rows, metrics = evaluate_predictions(classified, labels)
        classified = classified.merge(
            normalize_labels(labels), on="headline_id", how="left", validate="one_to_one"
        )
    daily = aggregate_daily_sentiment(classified)
    summary = company_summary(daily)
    outputs_dir = project_dir / "outputs"
    outputs_dir.mkdir(parents=True, exist_ok=True)
    classified.to_csv(outputs_dir / "classified_headlines.csv", index=False)
    daily.to_csv(outputs_dir / "daily_company_sentiment.csv", index=False)
    summary.to_csv(outputs_dir / "company_summary.csv", index=False)
    evaluation_rows.to_csv(outputs_dir / "evaluation_predictions.csv", index=False)
    with (outputs_dir / "model_metrics.json").open("w", encoding="utf-8") as handle:
        json.dump(metrics, handle, indent=2)
    chart_path = outputs_dir / "daily_sentiment_index.png"
    figure = make_sentiment_chart(daily, chart_path)
    plt.close(figure)
    ai_log = create_ai_use_log()
    ai_log.to_csv(project_dir / "ai_use_log.csv", index=False)
    _write_workbook(
        outputs_dir / "stock_news_sentiment_outputs.xlsx",
        classified, daily, summary, evaluation_rows, metrics, ai_log, chart_path,
    )
    results = {
        "classified": classified,
        "daily": daily,
        "summary": summary,
        "evaluation_rows": evaluation_rows,
        "metrics": metrics,
        "ai_log": ai_log,
        "recommendation": generate_recommendation(summary),
    }
    write_report(project_dir, results)
    return results


def main() -> None:
    project_dir = Path(__file__).resolve().parent.parent
    results = run_pipeline(project_dir)
    metrics = results["metrics"]
    if metrics:
        print(f"Evaluation rows: {metrics['evaluation_rows']}")
        print(f"Accuracy: {metrics['accuracy']:.2%}")
        print(f"Macro F1: {metrics['macro_f1']:.2%}")
    else:
        print("Gold labels not found; accuracy was not calculated.")
    print(f"Classified rows: {len(results['classified'])}")
    print(f"Outputs written to: {project_dir / 'outputs'}")


if __name__ == "__main__":
    main()
